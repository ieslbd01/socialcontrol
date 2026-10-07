"""Reports (PRD-10 / T-RPT-*): totals must reconcile with the attempt log."""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from socialcontrol.platforms.adapters.mock import MockAdapter
from socialcontrol.publisher.engine import run_once
from socialcontrol.reports import builder
from tests.integration.test_publisher import T0, World

pytestmark = pytest.mark.integration


@pytest.fixture
def w(clean_db):
    return World(clean_db)


def test_period_boundaries_use_dhaka_calendar():
    # 2026-10-01 20:00 UTC is 02:00 on 2026-10-02 in Dhaka
    ref = datetime(2026, 10, 1, 20, tzinfo=UTC)
    s, e = builder.period_for("daily", ref)
    assert s == datetime(2026, 10, 1, 18, tzinfo=UTC) and e - s == timedelta(days=1)
    ws, we = builder.period_for("weekly", datetime(2026, 10, 7, 6, tzinfo=UTC))  # Wednesday
    assert ws.astimezone(builder.DHAKA).weekday() == 0 and we - ws == timedelta(days=7)
    ms, me = builder.period_for("monthly", datetime(2026, 12, 15, tzinfo=UTC))
    assert ms.astimezone(builder.DHAKA).day == 1 and (me - ms).days == 31
    prev_s, prev_e = builder.previous_period("daily", datetime(2026, 10, 2, 3, tzinfo=UTC))
    assert prev_e == builder.period_for("daily", datetime(2026, 10, 2, 3, tzinfo=UTC))[0]
    with pytest.raises(ValueError):
        builder.period_for("hourly", ref)


def seed_activity(w):
    acc = w.account()
    q = w.queue(acc)
    ok1 = w.post(acc, q, slot_at=T0)
    ok2 = w.post(acc, q, slot_at=T0 + timedelta(minutes=1))
    bad = w.post(acc, q, slot_at=T0 + timedelta(minutes=2), caption="x" * 50)
    acc2 = w.account(mode="ASSISTED", state="NOT_CONFIGURED", short="wa")
    q2 = w.queue(acc2, name="WA")
    asst = w.post(acc2, q2, slot_at=T0 + timedelta(minutes=3))
    run_once(
        w.engine,
        lambda _s: MockAdapter({"caption_limit": 10}),
        T0 + timedelta(minutes=10),
        maintain=False,
    )
    return ok1, ok2, bad, asst


def window():
    return T0 - timedelta(hours=1), T0 + timedelta(days=1)


def test_report_totals_reconcile_with_attempt_log(w):
    seed_activity(w)
    s, e = window()
    with w.engine.connect() as c:
        rep = builder.build_report(c, "daily", s, e)
        log_ok = c.execute(
            text("select count(*) from publish_attempts where result in ('SUCCESS','CONFIRMED')")
        ).scalar_one()
        log_failed = c.execute(
            text("select count(*) from publish_attempts where result='FAILED'")
        ).scalar_one()
        delivered = c.execute(
            text("select count(*) from publish_attempts where result='DELIVERED'")
        ).scalar_one()
    sm = rep.summary
    assert sm["published"] == log_ok == 2
    assert sm["failed_attempts"] == log_failed == 1 and sm["final_failures"] == 1
    assert sm["assisted_delivered"] == delivered == 1 and sm["awaiting"] == 1
    assert sum(p["published"] for p in rep.by_platform) == 2
    assert [f["error_code"] for f in rep.failures] == ["E032"]


def test_empty_period_reports_zeros_not_errors(w):
    with w.engine.connect() as c:
        rep = builder.build_report(c, "monthly", T0 + timedelta(days=100), T0 + timedelta(days=130))
    assert rep.summary["published"] == 0 and rep.summary["final_failures"] == 0
    assert rep.by_platform == [] and rep.failures == []
    assert "Published: 0" in builder.to_text_summary(rep)


def test_exports_are_valid_and_consistent(w):
    seed_activity(w)
    s, e = window()
    with w.engine.connect() as c:
        rep = builder.build_report(c, "weekly", s, e)
    data = json.loads(builder.to_json(rep))
    assert data["summary"]["published"] == 2 and data["kind"] == "weekly"
    csv_text = builder.to_csv(rep)
    assert "summary,published,2" in csv_text and csv_text.startswith("section,key,value")
    wb = load_workbook(io.BytesIO(builder.to_xlsx(rep)))
    assert wb.sheetnames == ["Summary", "By platform", "By queue", "Failures", "Upcoming"]
    rows = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert rows["published"] == 2


def test_csv_export_neutralises_formula_injection(w):
    acc = w.account()
    pid = w.post(acc, w.queue(acc), slot_at=T0)
    run_once(
        w.engine,
        lambda _s: MockAdapter({"outcome": "validation"}),
        T0 + timedelta(minutes=1),
        maintain=False,
    )
    with w.engine.begin() as c:
        c.execute(
            text("update publish_attempts set error_message='=cmd|calc' where post_id=:p"),
            {"p": pid},
        )
    s, e = window()
    with w.engine.connect() as c:
        rep = builder.build_report(c, "daily", s, e)
    assert "'=cmd|calc" in builder.to_csv(rep) or "VALIDATION" in builder.to_csv(rep)
    assert ",=cmd" not in builder.to_csv(rep)


def test_save_report_persists_summary(w):
    seed_activity(w)
    s, e = window()
    with w.engine.begin() as c:
        rep = builder.build_report(c, "daily", s, e)
        rid = builder.save_report(c, rep)
        stored = c.execute(text("select kind, summary from reports where id=:i"), {"i": rid}).one()
    assert stored[0] == "daily" and stored[1]["published"] == 2


def test_text_summary_mentions_platforms(w):
    seed_activity(w)
    s, e = window()
    with w.engine.connect() as c:
        rep = builder.build_report(c, "daily", s, e)
    txt = builder.to_text_summary(rep)
    assert "Published: 2" in txt and "facebook_page" in txt
